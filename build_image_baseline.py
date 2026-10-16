"""Build the initial image baseline from images already uploaded to Shopify."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from src.images import image_filename, save_baseline


ROOT = Path(__file__).resolve().parent


def main() -> None:
    filenames: set[str] = set()
    for folder in (ROOT / "input" / "WAMI_Tafeln", ROOT / "input" / "WAMI_Zuschnitte"):
        for workbook_path in folder.glob("*.xlsx"):
            if ".before-remove-" in workbook_path.name:
                continue
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                try:
                    headers = list(next(rows))
                except StopIteration:
                    continue
                image_indexes = [
                    index for index, header in enumerate(headers)
                    if isinstance(header, str) and (header == "Bild" or header.startswith("Bild "))
                ]
                for row in rows:
                    for index in image_indexes:
                        if index < len(row):
                            filename = image_filename(row[index])
                            if filename:
                                filenames.add(filename)
            workbook.close()
    save_baseline(ROOT / "config" / "image_baseline.json", filenames)
    print(f"Saved {len(filenames)} image filenames to config/image_baseline.json")


if __name__ == "__main__":
    main()
