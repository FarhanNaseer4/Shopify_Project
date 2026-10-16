import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from config.categories import get_category
from src.converter import _current_rows, _write_products
from src.pricing import build_tafeln_price_data, build_zuschnitte_price_data


class CurrentWorkflowTests(unittest.TestCase):
    def test_existing_sku_keeps_shopify_handle_and_omits_media(self):
        headers = ["Handle", "Command", "Variant Command", "Variant SKU", "Variant Price",
                   "Image Src", "Unit Price Base Measure",
                   "Variant Metafield: custom.produktdatenblatt [file_reference]"]
        reference = [{"Handle": "shop-handle", "Variant SKU": "1", "Image Src": "old.jpg"}]
        updates, new = _current_rows([{"Artikel-Nr": 1, "Handle": "#N/A", "Preis 1": 12.345}],
                                    headers, get_category("Tafeln", "Lochbleche"), reference, [])
        self.assertEqual(new, [])
        self.assertEqual(updates[0]["Handle"], "shop-handle")
        self.assertEqual(updates[0]["Variant Price"], 12.35)
        self.assertEqual(updates[0]["Variant Command"], "UPDATE")
        self.assertFalse(any(h.startswith(("Image ", "Unit Price ")) or "produktdatenblatt" in h for h in updates[0]))

    def test_new_products_and_new_variants_use_distinct_product_commands(self):
        common = {"Material": "Stahl", "Stärke": "1 mm", "Format": "Kleinformat", "Preis 1": 10,
                  "Def-Lochung": "Test", "Lochgröße und -abstand": "T 1"}
        records = [{**common, "Artikel-Nr": 2, "Handle": "existing"},
                   {**common, "Artikel-Nr": 3, "Handle": "new-product"}]
        _, additions = _current_rows(records, ["Handle", "Variant SKU", "Variant Price"],
                                     get_category("Tafeln", "Lochbleche"),
                                     [{"Variant SKU": 1, "Handle": "existing"}], [])
        self.assertEqual([r["Command"] for r in additions], ["UPDATE", "NEW"])
        self.assertTrue(all(r["Variant Command"] == "NEW" for r in additions))
        self.assertEqual(additions[0]["Option3 Name"], "Format")
        self.assertEqual(additions[0]["Option3 Value"], "kleinformat")

    def test_output_removes_empty_columns_and_uses_products_sheet(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sample.xlsx"
            _write_products(path, [{"Handle": "a", "Variant SKU": "1", "Empty": None}])
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ["Products"])
            self.assertEqual([c.value for c in wb.active[1]], ["Handle", "Variant SKU"])
            wb.close()

    def test_rounding_and_suspicious_tier_warning(self):
        data, warnings = build_tafeln_price_data([{"Artikel-Nr": 1, "Preis 1": 200,
                                                 "Preis 2": 124.5888, "Preis 4": 130.005}])
        self.assertEqual(data["1"], {"2": 124.59, "4": 130.01})
        self.assertEqual(len(warnings), 1)

    def test_deleted_skus_cannot_reenter_price_json(self):
        data, _ = build_zuschnitte_price_data([{"Artikel-Nr": 1272, "Preis": 440},
                                             {"Artikel-Nr": 15838, "Preis": 15.5}])
        self.assertEqual(data, {"15838": 16})


if __name__ == "__main__":
    unittest.main()
