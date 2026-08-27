"""Safe, read-only workbook discovery and table loading."""

from __future__ import annotations

import unicodedata
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def normalized(value: str) -> str:
    return unicodedata.normalize("NFC", value).casefold()


def find_template(group: str, templates_dir: Path) -> Path:
    matches = [p for p in templates_dir.glob("*.xlsx") if normalized(group) in normalized(p.name)]
    if len(matches) != 1:
        raise ValueError(f"Expected exactly one template containing '{group}' in {templates_dir}; found {len(matches)}.")
    return matches[0]


def resolve_sheet(workbook: Any, wanted: str) -> str:
    matches = [name for name in workbook.sheetnames if normalized(name) == normalized(wanted)]
    if len(matches) != 1:
        raise ValueError(f"Worksheet '{wanted}' was not found uniquely. Available: {', '.join(workbook.sheetnames)}")
    return matches[0]


def extract_headers(template_file: Path, sheet: str) -> list[str]:
    wb = load_workbook(template_file, read_only=True, data_only=False)
    ws = wb[resolve_sheet(wb, sheet)]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    wb.close()
    if not headers or any(not isinstance(h, str) or not h for h in headers):
        raise ValueError(f"Template sheet '{sheet}' has invalid or blank headers.")
    # Inventory is maintained exclusively in WAMI and must never be imported
    # into Shopify, even if a future reference template contains this column.
    headers = [header for header in headers if header != "Inventory Quantity"]
    # The agency requires existing Shopify images to be replaced by images from
    # WAMI. Older supplied templates predate the Matrixify Image Command column.
    if "Image Command" not in headers:
        image_index = headers.index("Image Src")
        headers.insert(image_index + 1, "Image Command")
    return headers


def read_records(input_file: Path, sheet: str) -> list[dict[str, Any]]:
    wb = load_workbook(input_file, read_only=True, data_only=True)
    ws = wb[resolve_sheet(wb, sheet)]
    iterator = ws.iter_rows(values_only=True)
    try:
        raw_headers = next(iterator)
    except StopIteration as exc:
        raise ValueError(f"Input worksheet '{sheet}' is empty.") from exc
    last = max((i for i, value in enumerate(raw_headers) if value not in (None, "")), default=-1)
    headers = [str(value).strip() if value is not None else "" for value in raw_headers[: last + 1]]
    if len(set(headers)) != len(headers) or "" in headers:
        raise ValueError(f"Input worksheet '{sheet}' has duplicate or blank headers.")
    records = []
    for values in iterator:
        values = values[: len(headers)]
        if not any(value not in (None, "") for value in values):
            continue
        records.append(dict(zip(headers, values)))
    wb.close()
    if not records:
        raise ValueError(f"Input worksheet '{sheet}' contains no product rows.")
    return records
