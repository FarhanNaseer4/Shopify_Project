"""Generate all category samples and verify them against the current WAWI data."""
from collections import Counter
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
from config.categories import CATEGORIES
from src.converter import convert_wawi_to_shopify
from src.pricing import EXCLUDED_ZUSCHNITTE_SKUS, consolidated_filename, find_group_source
from src.workbook_reader import read_records
from src.transformer import sku_text

ROOT = Path(__file__).resolve().parent


def generate():
    before = {p: hashlib.sha256(p.read_bytes()).hexdigest()
              for p in (ROOT / 'input').rglob('*') if p.is_file()}
    out = ROOT / 'output' / ('client_samples_' + datetime.now().strftime('%Y%m%d_%H%M%S_%f'))
    results, sources = [], {}
    for group, categories in CATEGORIES.items():
        sources[group] = {}
        for category, config in categories.items():
            source = find_group_source(config, ROOT / 'input' / f'WAMI_{group}')
            records = read_records(source, config.source_sheet)
            expected = {sku_text(r['Artikel-Nr']): r for r in records
                        if group != 'Zuschnitte' or sku_text(r['Artikel-Nr']) not in EXCLUDED_ZUSCHNITTE_SKUS}
            sources[group].update(expected)
            result = convert_wawi_to_shopify(group=group, category=category, input_file=source,
                                             templates_dir=ROOT / 'input' / 'Templates',
                                             output_dir=out, logs_dir=ROOT / 'logs')
            results.append(result)
            seen = []
            for file in (result.output_file, result.new_products_file):
                if file is None:
                    continue
                wb = load_workbook(file, data_only=True)
                assert wb.sheetnames == ['Products'], file
                rows = list(wb.active.values)
                headers = rows[0]
                assert all(any(r[i] not in (None, '') for r in rows[1:]) for i in range(len(headers))), file
                assert not any(h.startswith(('Image ', 'Unit Price ')) or 'produktdatenblatt' in h
                               or h == 'Inventory Quantity' for h in headers), file
                assert not any(c.data_type == 'e' for row in wb.active for c in row), (file, 'Excel errors')
                previous_handles, last = set(), None
                for values in rows[1:]:
                    row = dict(zip(headers, values))
                    sku = sku_text(row['Variant SKU']); seen.append(sku)
                    assert row['Handle'] and not row['Handle'].startswith('#'), (file, sku)
                    if row['Handle'] != last:
                        assert row['Handle'] not in previous_handles, (file, 'noncontiguous product')
                        previous_handles.add(row['Handle']); last = row['Handle']
                    assert row['Command'] in {'UPDATE', 'NEW'}, (file, sku)
                    if group == 'Tafeln':
                        price = Decimal(str(expected[sku]['Preis 1'])).quantize(Decimal('.01'), rounding=ROUND_HALF_UP)
                        assert Decimal(str(row['Variant Price'])) == price, (file, sku)
                wb.close()
            assert Counter(seen) == Counter({sku: 1 for sku in expected}), (group, category, 'SKU coverage')
    for group, records in sources.items():
        data = json.loads((out / consolidated_filename(group)).read_text(encoding='utf-8'))
        assert set(data) == set(records), (group, 'JSON SKU coverage')
        for sku, source in records.items():
            if group == 'Tafeln':
                expected = {str(q): float(Decimal(str(source[f'Preis {q}'])).quantize(Decimal('.01'), rounding=ROUND_HALF_UP))
                            for q in (2, 4, 7, 10) if source.get(f'Preis {q}') not in (None, '')}
            else:
                expected = int(Decimal(str(source['Preis'])).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
            assert data[sku] == expected, (group, sku, 'JSON price')
    assert all(hashlib.sha256(p.read_bytes()).hexdigest() == digest for p, digest in before.items())
    lines = ['# Client sample review', '', 'No Shopify upload was performed. All input/template hashes remained unchanged.', '',
             '| Group | Category | Existing variants | New variants |', '|---|---|---:|---:|']
    for r in results:
        lines.append(f'| {r.group} | {r.category} | {r.variants-r.new_variants} | {r.new_variants} |')
    lines += ['', '## Files and workflow', '',
              'Each category update file has a Products sheet. Separate _NEW_ workbooks contain creation rows.',
              'Review creation files before import. Refresh Shopify reference templates after successful creation imports.',
              'Images, PDFs, Unit Price and inventory columns are omitted. Entirely empty columns are removed.',
              'Existing products receive template-mapped technical fields and Tafeln prices; titles/options are preserved.',
              'Complete JSON coverage: ' + ', '.join(f'{len(v)} {g} SKUs' for g,v in sources.items()) + '.',
              'No live products are deleted by these files.', '', '## Source warnings requiring review', '']
    lines.extend('- ' + w for w in sorted({w for r in results for w in r.warnings}))
    lines += ['', '### Increasing price tiers (source values)', '',
              '| SKU | Preis 1 | Preis 2 | Preis 4 | Preis 7 | Preis 10 |', '|---|---:|---:|---:|---:|---:|']
    for sku, source in sources['Tafeln'].items():
        prices = [source.get(f'Preis {q}') for q in (1, 2, 4, 7, 10)]
        if any(isinstance(a, (int, float)) and isinstance(b, (int, float)) and b > a for a,b in zip(prices, prices[1:])):
            formatted = [str(Decimal(str(p)).quantize(Decimal('.01'), rounding=ROUND_HALF_UP)) for p in prices]
            lines.append('| ' + ' | '.join([sku] + formatted) + ' |')
    lines += ['', '## Checks passed', '',
              '- Exact SKU coverage, no duplicates, and contiguous product rows in each category output.',
              '- Products sheet name; no empty columns, images/PDFs, inventory quantity or Unit Price fields.',
              '- Every Tafeln base price and both JSONs checked against WAWI using commercial rounding.',
              '- Confirmed obsolete SKUs and Prägeblech Tafelware excluded.',
              '', 'Local checks passed. Agency review and a Matrixify test import are still required.']
    (out / 'REVIEW_REPORT.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    print(out)
    print(f'Validated {len(results)} categories; {len(list(out.glob("*.xlsx")))} Excel files and 2 JSON files.')
    return out


if __name__ == '__main__':
    generate()
